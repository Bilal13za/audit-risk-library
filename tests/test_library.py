"""Tests for the 12 testing requirements in the project brief."""
import csv
import io
import json
import socket
import sys
from pathlib import Path

import pytest

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from db import connect, validate_db  # noqa: E402
from export import export_csv, export_excel  # noqa: E402


def active_rules(conn):
    rows = conn.execute("SELECT id, definition FROM rules WHERE status='Active'").fetchall()
    return [(r["id"], json.loads(r["definition"])) for r in rows]


def rule_all_factors(definition):
    return [c["factor_id"] for c in definition.get("when", {}).get("all", [])
            if c.get("type") == "factor_selected"]


# 1. Every predefined rule returns the expected risk -------------------------
def test_every_rule_returns_expected_risk(engine):
    for rule_id, definition in active_rules(engine.conn):
        factors = rule_all_factors(definition)
        if not factors:
            continue  # rules keyed on cycles only are covered separately
        plan = engine.assess(factors)
        got = {r["id"] for r in plan["risks"]}
        for rid in definition["then"].get("risk_ids", []):
            assert rid in got, f"{rule_id}: expected risk {rid} not returned"


# 2. Unrelated factors do not trigger irrelevant rules ------------------------
def test_unrelated_factors_do_not_trigger_irrelevant_rules(engine):
    all_factors = [r["id"] for r in
                   engine.conn.execute("SELECT id FROM business_factors WHERE review_status='Verified'")]
    for rule_id, definition in active_rules(engine.conn):
        referenced = set(rule_all_factors(definition))
        referenced |= {c.get("factor_id") for c in definition.get("when", {}).get("exclude", [])}
        referenced |= {c.get("factor_id") for c in definition.get("when", {}).get("any", [])}
        others = [f for f in all_factors if f not in referenced]
        assert others, "need at least one unrelated factor to test with"
        plan = engine.assess([others[0]])
        triggered = {t["rule_id"] for t in plan["triggered_rules"]}
        assert rule_id not in triggered, f"{rule_id} fired for unrelated factor {others[0]}"


# 3. Multiple factors return consolidated results -----------------------------
def test_multiple_factors_consolidated(engine):
    factors = [r["id"] for r in engine.conn.execute(
        "SELECT id FROM business_factors WHERE review_status='Verified' LIMIT 3")]
    assert len(factors) >= 2
    plan = engine.assess(factors)
    assert plan["counts"]["rules_triggered"] >= 1
    # determinism: same inputs -> same output
    plan2 = engine.assess(list(reversed(factors)))
    assert plan == plan2


# 4. Duplicate procedures are handled correctly -------------------------------
def test_duplicate_procedures_deduplicated(engine):
    factors = [r["id"] for r in engine.conn.execute(
        "SELECT id FROM business_factors WHERE review_status='Verified'")]
    plan = engine.assess(factors)
    consolidated_ids = [p["id"] for p in plan["procedures"]]
    assert len(consolidated_ids) == len(set(consolidated_ids)), "duplicate procedure ids in plan"
    per_risk_ids = {p["id"] for r in plan["risks"] for p in r["procedures"]}
    assert set(consolidated_ids) == per_risk_ids, "consolidated list does not cover all procedures"
    for p in plan["procedures"]:
        assert p["risk_ids"], f"procedure {p['id']} addresses no risk"


# 5. Missing information is reported ------------------------------------------
def test_missing_information_reported(engine):
    plan = engine.assess(["F-INV-001"])
    assert plan["missing_information"], "expected missing_information for RULE-INV-001"
    assert any("material" in m.lower() for m in plan["missing_information"])


# 6. Exclusion conditions work -------------------------------------------------
def test_exclusion_conditions(engine):
    # synthetic rule: fires on F-INV-001 unless F-REV-001 is also selected
    rule = {"when": {"all": [{"type": "factor_selected", "factor_id": "F-INV-001"}],
                        "any": [],
                        "exclude": [{"type": "factor_selected", "factor_id": "F-REV-001"}]},
            "then": {"risk_ids": ["R-INV-001"]}}
    matched, _ = engine._evaluate(rule, ["F-INV-001"], [])
    assert matched is True
    matched, _ = engine._evaluate(rule, ["F-INV-001", "F-REV-001"], [])
    assert matched is False
    # and a real seeded rule using exclude must exist (FRP management-override pack)
    n = engine.conn.execute(
        "SELECT COUNT(*) c FROM rules WHERE status='Active' AND definition LIKE '%exclude%'"
        " AND definition NOT LIKE '%\"exclude\": []%'").fetchone()["c"]
    assert n >= 1, "expected at least one seeded rule with a non-empty exclude list"


# 7. Evidence mappings are traceable to procedures -----------------------------
def test_evidence_traceable_to_procedures(engine):
    rows = engine.conn.execute(
        "SELECT pe.procedure_id, pe.evidence_id FROM procedure_evidence pe"
        " JOIN audit_procedures p ON p.id = pe.procedure_id"
        " JOIN evidence e ON e.id = pe.evidence_id"
        " WHERE p.review_status='Verified' AND e.review_status='Verified'").fetchall()
    assert rows, "no verified procedure-evidence links"
    plan = engine.assess(["F-INV-001"])
    for risk in plan["risks"]:
        for proc in risk["procedures"]:
            assert proc["evidence"], f"procedure {proc['id']} has no evidence in plan"


# 8. Rule versions are preserved ------------------------------------------------
def test_rule_versions_preserved(engine, db_path, tmp_path):
    from seed import load_seed_file
    conn = connect(db_path)
    current = conn.execute(
        "SELECT version, definition FROM rules WHERE id='RULE-INV-001' AND status='Active'").fetchone()
    definition = json.loads(current["definition"])
    bumped = {
        "rules": [{
            "rule_id": "RULE-INV-001", "name": "Inadequate warehouse security -> inventory loss risk",
            "version": current["version"] + 1, "priority": 10, "effective_date": "2026-10-22",
            "status": "Active", "review_status": "Verified", "reviewed_by": "test",
            "source_reference": "test", "rationale": "test version bump",
            "when": definition["when"], "then": definition["then"],
        }]
    }
    load_seed_file(conn, bumped, "test-bump", rules_dir=tmp_path)
    conn.commit()
    hist = conn.execute(
        "SELECT version, superseded_by FROM rule_history WHERE id='RULE-INV-001'").fetchall()
    assert hist and hist[0]["version"] == current["version"], "superseded version not archived"
    new_current = conn.execute(
        "SELECT version FROM rules WHERE id='RULE-INV-001' AND status='Active'").fetchone()
    assert new_current["version"] == current["version"] + 1
    conn.close()


# 9. Exported workpapers contain the selected data ------------------------------
def test_export_contains_selected_data(engine, tmp_path):
    plan = engine.assess(["F-INV-001"])
    csv_path = tmp_path / "wp.csv"
    n = export_csv(plan, csv_path)
    assert n > 0
    text = csv_path.read_text(encoding="utf-8")
    for token in ("R-INV-001", "P-INV-001", "BAL-EXI", "Expected result", "Follow-up actions"):
        assert token in text, f"CSV missing {token}"
    rows = list(csv.DictReader(io.StringIO(text)))
    assert all(r["Risk ID"] == "R-INV-001" for r in rows)

    xlsx_path = tmp_path / "wp.xlsx"
    export_excel(plan, xlsx_path)
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    assert set(wb.sheetnames) >= {"Cover", "Risk register", "Audit procedures", "Evidence", "Rules applied"}
    ws = wb["Risk register"]
    values = [c.value for row in ws.iter_rows(min_row=2) for c in row]
    assert "R-INV-001" in values


# 10. The application operates offline ------------------------------------------
def test_operates_offline(engine, tmp_path, monkeypatch):
    def _blocked(*a, **k):
        raise AssertionError("network access attempted")
    monkeypatch.setattr(socket, "socket", _blocked)
    monkeypatch.setattr(socket, "create_connection", _blocked)
    plan = engine.assess(["F-INV-001", "F-REV-001"])
    assert plan["risks"]
    assert engine.search("warehouse")
    export_csv(plan, tmp_path / "off.csv")
    export_excel(plan, tmp_path / "off.xlsx")


# 11. Invalid or incomplete database records are detected -----------------------
def test_invalid_records_detected(engine):
    conn = engine.conn
    conn.execute(
        "INSERT INTO risk_events (id, title, description, cycle_id, review_status)"
        " VALUES ('R-BAD-001', 'Bad risk', 'no procedures linked', 'CYC-INV', 'Verified')")
    try:
        errors = validate_db(conn)
        assert any("R-BAD-001" in e for e in errors), "validator missed orphan verified risk"
    finally:
        conn.execute("DELETE FROM risk_events WHERE id='R-BAD-001'")
        conn.commit()


# 12. Every verified mapping passes validation ----------------------------------
def test_verified_mappings_pass_validation(engine):
    errors = validate_db(engine.conn)
    assert errors == [], f"validation errors:\n" + "\n".join(errors)


# End-to-end acceptance test: the inventory theft example from the brief --------
def test_inventory_theft_end_to_end(engine):
    plan = engine.assess(["F-INV-001"])

    assert [t["rule_id"] for t in plan["triggered_rules"]] == ["RULE-INV-001"]
    rule = plan["triggered_rules"][0]
    assert rule["version"] >= 1 and rule["rationale"]

    assert len(plan["risks"]) == 1
    risk = plan["risks"][0]
    assert risk["id"] == "R-INV-001"
    assert "overstated" in (risk["potential_consequences"] or "").lower()

    assertions = {a["id"]: a for a in risk["assertions"]}
    assert set(assertions) == {"BAL-EXI", "BAL-COM", "BAL-VAL"}
    for a in assertions.values():
        assert a["why_relevant"], f"assertion {a['id']} has no rationale"

    proc_ids = [p["id"] for p in risk["procedures"]]
    assert proc_ids == [f"P-INV-00{i}" for i in range(1, 9)]
    for p in risk["procedures"]:
        assert p["expected_result"], f"{p['id']} missing expected result"
        assert p["exception_criteria"], f"{p['id']} missing exception criteria"
        assert p["follow_up"], f"{p['id']} missing follow-up"
        assert "fraud has occurred" not in p["follow_up"].lower() or "not" in p["follow_up"].lower()
        assert p["evidence"], f"{p['id']} has no evidence"
        for e in p["evidence"]:
            assert e["reliability_notes"], f"evidence {e['id']} missing reliability notes"

    assert len(risk["controls"]) >= 1
    assert plan["missing_information"]


# ---- v1.2 feature tests: industries, fraud, analytical procedures, controls --

def test_industry_factors_included_in_assessment(engine):
    plan = engine.assess([], industry_ids=["IND-TEL"])
    risk_ids = {r["id"] for r in plan["risks"]}
    assert "R-TEL-001" in risk_ids, "telecom decommissioning risk not raised by industry selection"
    assert plan["industry_added_factors"], "industry-added factors not reported"
    assert all(f["industry_id"] == "IND-TEL" for f in plan["industry_added_factors"])


def test_industry_assessment_deterministic(engine):
    p1 = engine.assess(["F-INV-001"], industry_ids=["IND-MFG"])
    p2 = engine.assess(["F-INV-001"], industry_ids=["IND-MFG"])
    assert p1 == p2


def test_fraud_indicator_triggers_fraud_responses(engine):
    plan = engine.assess(["F-FRD-001"])
    assert plan["fraud_responses"], "no fraud responses for segregation-of-duties indicator"
    assert all(p["fraud_focus"] for p in plan["fraud_responses"])
    assert any("R-FRD" in r["id"] for r in plan["risks"])


def test_significant_risks_flagged(engine):
    plan = engine.assess(["F-FRD-002"])
    assert plan["significant_risks"], "significant risks not flagged"
    for risk in plan["risks"]:
        if risk["id"] in plan["significant_risks"]:
            assert risk["significance"] == "Significant risk"


def test_fraud_rules_run_first(engine):
    rows = engine.conn.execute(
        "SELECT id, priority FROM rules WHERE status='Active' AND id LIKE 'RULE-FRD-%'").fetchall()
    assert rows, "no fraud rules seeded"
    assert all(r["priority"] == 5 for r in rows)


def test_fraud_triangle_values_valid(engine):
    rows = engine.conn.execute(
        "SELECT DISTINCT fraud_triangle FROM business_factors"
        " WHERE risk_category='Fraud risk' AND fraud_triangle IS NOT NULL").fetchall()
    valid = {"Incentive or pressure", "Opportunity", "Rationalization"}
    assert {r["fraud_triangle"] for r in rows} <= valid


def test_analytical_procedures_by_area(engine):
    areas = engine.analytical_areas()
    assert len(areas) >= 7
    rev = engine.analytical_procedures("Revenue")
    assert len(rev) == 3
    for ap in rev:
        assert ap["detail"] and ap["expected_relationship"]
        assert ap["investigation_threshold"] and ap["follow_up"]


def test_control_objectives_bidirectional(engine):
    objectives = engine.control_objectives()
    assert len(objectives) == 12
    for o in objectives:
        ctrls = engine.controls_for_objective(o["id"])
        assert ctrls, f"objective {o['id']} has no controls"
    # reverse direction resolves back
    first_ctrl = engine.controls_for_objective(objectives[0]["id"])[0]["id"]
    back = [o["id"] for o in engine.objectives_for_control(first_ctrl)]
    assert objectives[0]["id"] in back


def test_analytical_export_unaffected(engine, tmp_path):
    plan = engine.assess(["F-TEL-001"], industry_ids=["IND-TEL"])
    n = export_csv(plan, tmp_path / "tel.csv")
    assert n > 0
    text = (tmp_path / "tel.csv").read_text(encoding="utf-8")
    assert "R-TEL-001" in text
